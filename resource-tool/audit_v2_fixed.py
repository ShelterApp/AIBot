import re,sys,time,argparse,phonenumbers
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill,Border,Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

URL="https://shelter.app/manage_services"
# abc@gmail.com is NOT a placeholder in ShelterApp data: never flag or clear it.
PLACEHOLDER_EMAILS=["test@test.com","email@email.com","test@example.com","example@example.com"]
PLACEHOLDER_WEBSITES=["https://your-nonprofit.org","http://your-nonprofit.org","your-nonprofit.org","https://www.example.com","http://example.com"]
PLACEHOLDER_SOCIAL=["https://www.facebook.com/abc.xyz","https://www.twitter.com/abc.xyz","https://www.instagram.com/abc.xyz"]
TEST_NAMES=["test service","test resource","test shelter","test","sample","demo","placeholder","asdf"]

def find_input_by_label(page, label_text):
    for strategy in [
        lambda: page.locator(f"text='{label_text}'").first.locator("xpath=..").locator("input, textarea").first,
        lambda: page.locator(f"text='{label_text}'").first.locator("xpath=../..").locator("input, textarea").first,
    ]:
        try:
            el = strategy()
            if el.is_visible(timeout=1000): return el
        except: pass
    return None

def read_input(page, label_text):
    el = find_input_by_label(page, label_text)
    if el:
        try: return el.input_value().strip()
        except: pass
    return ""

def write_input(page, label_text, value):
    el = find_input_by_label(page, label_text)
    if el:
        try: el.fill(value); return True
        except: pass
    return False

def read_checkbox(page, label_text):
    try:
        cb = page.locator(f"text='{label_text}'").first.locator("xpath=..").locator("input[type='checkbox']").first
        return cb.is_checked(timeout=1000)
    except:
        try:
            cb = page.locator(f"text='{label_text}'").first.locator("xpath=../..").locator("input[type='checkbox']").first
            return cb.is_checked(timeout=1000)
        except: return None

def read_all_fields(page):
    d = {}
    print("    Reading fields via JS...", end="", flush=True)

    raw = page.evaluate("""() => {
        function getLabelText(inp) {
            // 1. aria-label attribute
            let t = inp.getAttribute('aria-label');
            if (t) return t.trim();
            // 2. <label for="id">
            if (inp.id) {
                const lbl = document.querySelector('label[for="' + inp.id + '"]');
                if (lbl) return lbl.textContent.trim();
            }
            // 3. Input wrapped inside a <label>
            const parentLbl = inp.closest('label');
            if (parentLbl) {
                return Array.from(parentLbl.childNodes)
                    .filter(n => n !== inp && !n.contains(inp))
                    .map(n => n.textContent.trim()).filter(Boolean).join(' ');
            }
            // 4. Walk up to 4 levels; at each level scan preceding siblings for a text element
            let el = inp;
            for (let depth = 0; depth < 4; depth++) {
                const parent = el.parentElement;
                if (!parent) break;
                const siblings = Array.from(parent.childNodes);
                const idx = siblings.indexOf(el);
                for (let i = idx - 1; i >= 0; i--) {
                    const sib = siblings[i];
                    if (sib.nodeType === 1) {
                        const txt = sib.textContent.trim();
                        if (txt && txt.length < 100 && !sib.querySelector('input, textarea, select'))
                            return txt;
                    } else if (sib.nodeType === 3) {
                        const txt = sib.textContent.trim();
                        if (txt) return txt;
                    }
                }
                el = parent;
            }
            // 5. Fallback: placeholder / name attribute
            return inp.getAttribute('placeholder') || inp.getAttribute('name') || '';
        }

        const result = {};
        for (const inp of document.querySelectorAll('input, textarea')) {
            if (['hidden','submit','button','image','reset'].includes(inp.type)) continue;
            const label = getLabelText(inp).toUpperCase();
            if (!label) continue;
            if (inp.type === 'checkbox' || inp.type === 'radio') {
                result[label] = inp.checked;
            } else {
                result[label] = inp.value.trim();
            }
        }
        return result;
    }""")

    def _lookup(label):
        # exact label first, then substring (labels like
        # "COMMUNITY SERVICE/ VOLUNTEER NAME" combine two names)
        lu = label.upper()
        if lu in raw:
            return raw[lu]
        for k, v in raw.items():
            if lu in k:
                return v
        return None

    def get_field(*labels):
        for lbl in labels:
            v = _lookup(lbl)
            if v is not None and v != '': return v
        return ''

    def get_bool(*labels):
        for lbl in labels:
            v = _lookup(lbl)
            if v is not None: return bool(v)
        return None

    d['name']            = get_field('COMMUNITY SERVICE', 'VOLUNTEER NAME')
    d['description']     = get_field('DESCRIPTION')
    d['address1']        = get_field('ADDRESS 1')
    d['address2']        = get_field('ADDRESS 2')
    d['city']            = get_field('CITY')
    d['state']           = get_field('STATE')
    d['zip']             = get_field('ZIP')
    d['phone']           = get_field('PHONE')
    d['email']           = get_field('EMAIL')
    d['website']         = get_field('WEBSITE')
    d['facebook']        = get_field('FACEBOOK')
    d['twitter']         = get_field('TWITTER')
    d['service_details'] = get_field('SERVICE DETAILS')
    d['age_group']       = get_field('AGE GROUP')
    d['owned_by']        = get_field('OWNED BY')

    # Debug: if name is still blank, show what labels JS found so you can tune get_field() calls
    if not d['name'] and raw:
        text_keys = [k for k, v in raw.items() if isinstance(v, str) and v]
        print(f"\n    [DEBUG] JS found text fields: {text_keys[:20]}", end="")

    print(" checkboxes...", end="", flush=True)
    services = []
    for svc in ['FOOD', 'SHELTER', 'HEALTH', 'RESOURCES', 'WORK']:
        if get_bool(svc): services.append(svc)
    d['services'] = ', '.join(services)

    categories = []
    for cat in ['MEN', 'WOMEN', 'YOUTH/KIDS', 'SENIORS', 'DISABLED', 'FAMILIES', 'LGBT+', 'ALL']:
        if get_bool(cat): categories.append(cat)
    d['categories'] = ', '.join(categories)

    d['contact_for_hours'] = get_bool('PLEASE CONTACT THIS SERVICE FOR HOURS')
    d['never_expires']     = get_bool('NEVER EXPIRES')

    schedule_entries = []
    try:
        entries = page.locator("text=/MONDAY|TUESDAY|WEDNESDAY|THURSDAY|FRIDAY|SATURDAY|SUNDAY/i").all()
        for entry in entries:
            try:
                row_text = entry.locator("xpath=..").inner_text(timeout=1000).strip()
                if row_text and ("AM" in row_text or "PM" in row_text):
                    schedule_entries.append(row_text.replace("\n", " ").strip())
            except: pass
    except: pass
    d['schedule'] = '; '.join(schedule_entries) if schedule_entries else ''

    print(" done")
    return d

def validate_phone(p):
    if not p: return False, ["Missing phone"]
    try:
        pn = phonenumbers.parse(p, "US")
        if not phonenumbers.is_valid_number(pn): return False, [f"Invalid phone: {p}"]
        return True, []
    except: return False, [f"Cannot parse phone: {p}"]

def validate_email(e):
    if not e: return False, ["Missing email"]
    if e.lower() in PLACEHOLDER_EMAILS: return False, [f"Placeholder email: {e}"]
    if not re.match(r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$', e):
        return False, [f"Invalid email format: {e}"]
    return True, []

def validate_website(w):
    if not w: return False, ["Missing website"]
    if w.lower().strip().rstrip("/") in [x.rstrip("/") for x in PLACEHOLDER_WEBSITES]:
        return False, [f"Placeholder website: {w}"]
    url = w if w.startswith("http") else "https://"+w
    try:
        parsed = urlparse(url)
        if parsed.scheme and parsed.netloc and "." in parsed.netloc: return True, []
    except: pass
    return False, [f"Invalid URL: {w}"]

def is_placeholder_social(url):
    return url.lower().strip().rstrip("/") in [x.rstrip("/") for x in PLACEHOLDER_SOCIAL]

def audit_resource(d):
    issues = []; changes = []
    name = d.get("name","")
    if not name: issues.append(("HIGH","MISSING: Name"))
    elif name.lower().strip() in TEST_NAMES: issues.append(("HIGH",f"TEST DATA: Name is '{name}'"))
    desc = d.get("description","")
    if not desc: issues.append(("MEDIUM","MISSING: Description"))
    elif desc.lower().strip() in TEST_NAMES: issues.append(("HIGH",f"TEST DATA: Description is '{desc}'"))
    if not d.get("address1"): issues.append(("HIGH","MISSING: Address"))
    if not d.get("city"): issues.append(("HIGH","MISSING: City"))
    if not d.get("state"): issues.append(("HIGH","MISSING: State"))
    elif len(d["state"])!=2: issues.append(("MEDIUM",f"STATE: Not 2-letter: '{d['state']}'"))
    if not d.get("zip"): issues.append(("MEDIUM","MISSING: ZIP code"))
    elif not re.match(r"^\d{5}(-\d{4})?$",d["zip"]): issues.append(("MEDIUM",f"ZIP: Invalid '{d['zip']}'"))
    valid, pi = validate_phone(d.get("phone"))
    if not valid:
        for x in pi: issues.append(("HIGH" if "Missing" in x else "MEDIUM",f"PHONE: {x}"))
    email = d.get("email","")
    valid, ei = validate_email(email)
    if not valid:
        for x in ei: issues.append(("HIGH" if "Placeholder" in x else "MEDIUM",f"EMAIL: {x}"))
    if email.lower() in PLACEHOLDER_EMAILS:
        changes.append(("email",email,"","Cleared placeholder email"))
    website = d.get("website","")
    valid, wi = validate_website(website)
    if not valid:
        for x in wi: issues.append(("HIGH" if "Placeholder" in x else "MEDIUM",f"WEBSITE: {x}"))
    if website.lower().strip().rstrip("/") in [x.rstrip("/") for x in PLACEHOLDER_WEBSITES]:
        changes.append(("website",website,"","Cleared placeholder website"))
    fb = d.get("facebook","")
    if fb and is_placeholder_social(fb):
        issues.append(("MEDIUM","FACEBOOK: Placeholder URL"))
        changes.append(("facebook",fb,"","Cleared placeholder Facebook"))
    tw = d.get("twitter","")
    if tw and is_placeholder_social(tw):
        issues.append(("MEDIUM","TWITTER: Placeholder URL"))
        changes.append(("twitter",tw,"","Cleared placeholder Twitter"))
    if not d.get("services"): issues.append(("MEDIUM","MISSING: No service type selected"))
    sd = d.get("service_details","")
    if sd and sd.lower().strip() in TEST_NAMES:
        issues.append(("HIGH",f"TEST DATA: Service Details is '{sd}'"))
        changes.append(("service_details",sd,"","Cleared test service details"))
    if not d.get("categories"): issues.append(("MEDIUM","MISSING: No categories selected"))
    if not d.get("schedule") and not d.get("contact_for_hours"):
        issues.append(("MEDIUM","MISSING: No schedule"))
    owned = d.get("owned_by","")
    if owned.lower() in PLACEHOLDER_EMAILS:
        issues.append(("HIGH",f"OWNED BY: Placeholder ({owned})"))
    sevs = [i[0] for i in issues]
    overall = "HIGH" if "HIGH" in sevs else ("MEDIUM" if "MEDIUM" in sevs else ("LOW" if issues else "PASS"))
    return issues, changes, overall

def apply_corrections(page, changes, dry_run=False):
    applied = []
    label_map = {"email":"EMAIL","website":"WEBSITE","facebook":"FACEBOOK","twitter":"TWITTER","service_details":"SERVICE DETAILS"}
    for field, old_val, new_val, reason in changes:
        if dry_run: applied.append((field,old_val,new_val,reason,"DRY RUN")); continue
        label = label_map.get(field)
        if label:
            if write_input(page, label, new_val): applied.append((field,old_val,new_val,reason,"APPLIED"))
            else: applied.append((field,old_val,new_val,reason,"FAILED"))
    return applied

def export_xlsx(results, output_path):
    wb = Workbook()
    hf = Font(name="Arial",bold=True,color="FFFFFF",size=11)
    hfill = PatternFill("solid",fgColor="2F5496")
    ha = Alignment(horizontal="center",vertical="center",wrap_text=True)
    bd = Border(left=Side("thin"),right=Side("thin"),top=Side("thin"),bottom=Side("thin"))
    sf = {"HIGH":PatternFill("solid",fgColor="FFC7CE"),"MEDIUM":PatternFill("solid",fgColor="FFEB9C"),
          "LOW":PatternFill("solid",fgColor="C6EFCE"),"PASS":PatternFill("solid",fgColor="FFFFFF")}
    ws = wb.active; ws.title = "Audit Findings"
    hdrs = ["#","Severity","Name","Address 1","City","State","ZIP","Phone","Email","Website",
            "Services","Categories","Schedule","Description","Service Details","Owned By","Issues","Count"]
    for c,h in enumerate(hdrs,1):
        cl=ws.cell(row=1,column=c,value=h); cl.font=hf; cl.fill=hfill; cl.alignment=ha; cl.border=bd
    for i,r in enumerate(results,2):
        d=r["data"]
        it="\n".join(f"[{s}] {t}" for s,t in r["issues"]) if r["issues"] else "No issues"
        row=[i-1,r["severity"],d.get("name",""),d.get("address1",""),d.get("city",""),d.get("state",""),
             d.get("zip",""),d.get("phone",""),d.get("email",""),d.get("website",""),d.get("services",""),
             d.get("categories",""),d.get("schedule",""),d.get("description","")[:100],
             d.get("service_details",""),d.get("owned_by",""),it,len(r["issues"])]
        for c,v in enumerate(row,1):
            cl=ws.cell(row=i,column=c,value=v); cl.border=bd; cl.alignment=Alignment(vertical="top",wrap_text=True)
            if c==2: cl.fill=sf.get(r["severity"],PatternFill()); cl.font=Font(bold=True)
    widths=[5,9,30,25,14,6,8,16,22,26,20,28,28,25,20,22,42,6]
    for c,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes="A2"
    ws.auto_filter.ref=f"A1:{get_column_letter(len(hdrs))}{len(results)+1}"
    ws2=wb.create_sheet("Changes Made")
    for c,h in enumerate(["#","Resource","Field","Old","New","Reason","Status"],1):
        cl=ws2.cell(row=1,column=c,value=h); cl.font=hf; cl.fill=hfill; cl.border=bd
    rn=2
    for r in results:
        for ch in r.get("applied",[]):
            f,o,n,reason,st=ch
            ws2.cell(row=rn,column=1,value=rn-1); ws2.cell(row=rn,column=2,value=r["data"].get("name",""))
            ws2.cell(row=rn,column=3,value=f); ws2.cell(row=rn,column=4,value=o)
            ws2.cell(row=rn,column=5,value=n if n else "(cleared)"); ws2.cell(row=rn,column=6,value=reason)
            ws2.cell(row=rn,column=7,value=st); rn+=1
    ws3=wb.create_sheet("Summary")
    t=len(results); hi=sum(1 for r in results if r["severity"]=="HIGH")
    me=sum(1 for r in results if r["severity"]=="MEDIUM"); pa=sum(1 for r in results if r["severity"]=="PASS")
    tc=sum(len(r.get("applied",[])) for r in results)
    for row in [["ShelterApp Audit+Fix Summary"],[f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"],
                [""],["Metric","Count"],["Total",t],["HIGH",hi],["MEDIUM",me],["PASS",pa],["Corrections",tc],
                [""],["Common Issues","Count"]]:
        ws3.append(row)
    ic={}
    for r in results:
        for s,issue in r["issues"]: k=issue.split(":")[0]; ic[k]=ic.get(k,0)+1
    for k,v in sorted(ic.items(),key=lambda x:-x[1]): ws3.append([k,v])
    ws3["A1"].font=Font(bold=True,size=14); ws3.column_dimensions["A"].width=30
    ws4=wb.create_sheet("Flagged for Review")
    for c,h in enumerate(["#","Name","Reason","Owned By"],1):
        cl=ws4.cell(row=1,column=c,value=h); cl.font=hf; cl.fill=hfill
    fl=[r for r in results if any("TEST DATA" in i[1] for i in r["issues"])]
    for i,r in enumerate(fl,2):
        d=r["data"]; reasons=[x[1] for x in r["issues"] if "TEST DATA" in x[1]]
        ws4.cell(row=i,column=1,value=i-1); ws4.cell(row=i,column=2,value=d.get("name",""))
        ws4.cell(row=i,column=3,value="; ".join(reasons)); ws4.cell(row=i,column=4,value=d.get("owned_by",""))
    wb.save(output_path); print(f"\nSaved: {output_path}")

def get_edit_buttons(pg):
    buttons = []
    try:
        all_btns = pg.locator("text=Edit Service").all()
        for btn in all_btns:
            try:
                if not btn.is_visible(timeout=500): continue
                par = btn
                for _ in range(8):
                    par = par.locator("xpath=..")
                    try:
                        txt = par.inner_text(timeout=1000)
                        if "Phone:" in txt or "Address:" in txt or "Services:" in txt:
                            lines = [l.strip() for l in txt.split("\n") if l.strip()]
                            name = ""
                            for l in lines:
                                if l not in ("...","Edit Service","Delete Service","","This Service is waiting for Approval","Manage Services"):
                                    name = l
                                    break
                            if name: buttons.append({"name": name, "button": btn, "text": txt})
                            break
                    except: continue
            except: continue
    except: pass
    return buttons

def run(args):
    Path("audit_results").mkdir(exist_ok=True)
    with sync_playwright() as p:
        print("Launching browser...\n*** DO NOT CLOSE the browser window! ***\n")
        ud = str(Path.home()/".shelterapp-audit-browser")
        ctx = p.chromium.launch_persistent_context(ud,headless=False,slow_mo=100,
            viewport={"width":1400,"height":900},
            args=["--disable-blink-features=AutomationControlled"])
        pg = ctx.pages[0] if ctx.pages else ctx.new_page()
        url = args.url or URL
        print(f"Going to: {url}")
        try: pg.goto(url,wait_until="domcontentloaded",timeout=30000)
        except Exception as e: print(f"Note: {e}")
        print("\n"+"="*55)
        print("  Log in, get to Manage Services,")
        print("  then come back here and press ENTER.")
        print("  *** Keep browser OPEN! ***")
        print("="*55)
        input("\n>>> Press ENTER after login... ")
        ap = ctx.pages; pg = ap[-1] if len(ap)>1 else pg
        try: pg.wait_for_timeout(2000)
        except: pass
        list_url = pg.url
        mode = "FIX" if args.fix else "AUDIT"
        dry = args.dry_run
        if mode=="FIX" and dry: print("\nDRY RUN - showing changes without applying\n")
        elif mode=="FIX": print("\nFIX MODE - will auto-correct placeholder data\n")
        else: print("\nAUDIT MODE - read-only\n")
        print("Scrolling to find all resources...")
        seen = set(); resource_list = []
        no_new = 0; mx = args.max or 999999
        for scroll in range(200):
            btns = get_edit_buttons(pg)
            nf = 0
            for b in btns:
                if b["name"] not in seen and len(resource_list) < mx:
                    seen.add(b["name"]); resource_list.append(b["name"]); nf += 1
                    print(f"  Found: {b['name'][:60]}")
            if nf == 0: no_new += 1
            else: no_new = 0
            if no_new >= 5: break
            try: pg.evaluate("window.scrollBy(0, 800)"); pg.wait_for_timeout(1500)
            except: break
        total = len(resource_list)
        print(f"\nFound {total} resources. Auditing each one...\n")
        results = []
        for idx, rname in enumerate(resource_list):
            print(f"\n[{idx+1}/{total}] {rname[:55]}")
            try: pg.goto(list_url, wait_until="domcontentloaded", timeout=15000); pg.wait_for_timeout(2000)
            except: print("  ERROR: Could not load list page"); continue
            clicked = False
            for scroll in range(30):
                btns = get_edit_buttons(pg)
                for b in btns:
                    if b["name"] == rname:
                        try:
                            b["button"].scroll_into_view_if_needed(timeout=3000)
                            b["button"].click(timeout=5000); clicked = True
                        except: pass
                        break
                if clicked: break
                try: pg.evaluate("window.scrollBy(0, 600)"); pg.wait_for_timeout(800)
                except: break
            if not clicked: print("  Could not find/click Edit - skipping"); continue
            try: pg.wait_for_timeout(2500)
            except: pass
            try: pg.locator("text=Edit Service").first.inner_text(timeout=3000)
            except: print("  Edit page did not load - skipping"); continue
            data = read_all_fields(pg)
            if not data.get("name"): data["name"] = rname
            if args.state and data.get("state","").upper() != args.state.upper():
                print(f"  Skipping (state: {data.get('state','')})")
                continue
            issues, changes, severity = audit_resource(data)
            ic = {"HIGH":"!!","MEDIUM":"! ","LOW":"~ ","PASS":"OK"}[severity]
            print(f"  [{ic}] {len(issues)} issues")
            for sev, issue in issues[:5]: print(f"       [{sev}] {issue}")
            if len(issues) > 5: print(f"       ... +{len(issues)-5} more")
            applied = []
            if mode == "FIX" and changes:
                applied = apply_corrections(pg, changes, dry_run=dry)
                for field,old,new,reason,status in applied:
                    nd = new if new else "(cleared)"
                    print(f"       {status}: {field}: '{old}' -> '{nd}'")
                if applied and not dry:
                    try:
                        pg.locator("text=Update Service").first.click(timeout=5000)
                        pg.wait_for_timeout(2000); print("       Saved!")
                    except: print("       WARNING: Could not click Update Service")
            results.append({"data":data,"issues":issues,"severity":severity,"applied":applied})
        if results:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            ss = f"_{args.state.upper()}" if args.state else ""
            mm = "_fix" if mode=="FIX" else ""
            dr = "_dryrun" if dry else ""
            op = f"audit_results/shelterapp_audit{ss}{mm}{dr}_{ts}.xlsx"
            export_xlsx(results, op)
            print("\n"+"="*55)
            print(f"DONE - {len(results)} resources audited")
            hi=sum(1 for r in results if r["severity"]=="HIGH")
            me=sum(1 for r in results if r["severity"]=="MEDIUM")
            pa=sum(1 for r in results if r["severity"]=="PASS")
            tc=sum(len(r.get("applied",[])) for r in results)
            print(f"  HIGH: {hi}  MEDIUM: {me}  PASS: {pa}")
            print(f"  Auto-corrections: {tc}")
            print(f"  File: {op}")
        else: print("\nNo resources audited.")
        print("\nClosing browser..."); ctx.close()

if __name__=="__main__":
    pr=argparse.ArgumentParser()
    pr.add_argument("--state",type=str,default=None)
    pr.add_argument("--max",type=int,default=None)
    pr.add_argument("--url",type=str,default=None)
    pr.add_argument("--fix",action="store_true")
    pr.add_argument("--dry-run",action="store_true")
    run(pr.parse_args())
