"""
triage.py — Sort the flagged listings into actionable buckets (no network).

Reads the latest audit report + master inventory and categorizes HIGH/MEDIUM
listings so junk can be deleted fast and genuine ones spotted. Writes
audit_results/triage.csv (sorted by action).

    py triage.py
"""
import glob, os, json, csv, re
from openpyxl import load_workbook
import match_audit as ma

TEST_RE = re.compile(r"\b(test|sample|demo|placeholder|xxx|raj)\b", re.I)
# domains that look made-up rather than real org sites
FAKE_HINTS = ("balldaws", "baop.org", "bchomeless", "autaugacountyhomeless",
              "sdcpo", "sdhsa", "projectbreadline")


def fake_domain(site):
    s = (site or "").lower()
    return any(h in s for h in FAKE_HINTS)


def main():
    # newest audit report, whether written by run.py (audit_*.xlsx) or the
    # older pipeline flow (pipeline_*.xlsx)
    reports = glob.glob("audit_results/audit_*.xlsx") + glob.glob("audit_results/pipeline_*.xlsx")
    f = sorted(reports, key=os.path.getmtime)[-1]
    ws = load_workbook(f)["Audit Findings"]
    report = [(r[2], r[1], r[3], r[8]) for r in ws.iter_rows(min_row=2, values_only=True)]
    master = {d["name"]: d for d in json.load(open("audit_results/inventory_master.json", encoding="utf-8"))}
    shared = ma.shared_address_map(list(master.values()))

    def categorize(name, verdict, addr, finding):
        site = master.get(name, {}).get("website", "")
        n = shared.get(ma.norm_addr(addr), 1)
        if TEST_RE.search(name or ""):
            return "DELETE", "test/sample entry"
        if fake_domain(site):
            return "DELETE", "fabricated website (seed data)"
        if n >= 3:
            return "VERIFY-or-DELETE", f"placeholder address shared by {n} listings"
        if "UNRESOLVED" in (finding or ""):
            return "MANUAL-CHECK", "address+org not found online"
        if n == 2:
            return "VERIFY", "address shared with 1 other listing"
        if "MISMATCH" in (finding or "") or "did not geocode" in (finding or ""):
            return "MANUAL-CHECK", "address needs verification"
        return "REVIEW", (finding or "")[:50]

    order = {"DELETE": 0, "VERIFY-or-DELETE": 1, "MANUAL-CHECK": 2, "VERIFY": 3, "REVIEW": 4}
    out = []
    for name, verdict, addr, finding in report:
        if verdict not in ("HIGH", "MEDIUM"):
            continue
        action, reason = categorize(name, verdict, addr, finding)
        d = master.get(name, {})
        out.append({"action": action, "reason": reason, "name": name, "verdict": verdict,
                    "address": addr or "", "website": d.get("website", ""),
                    "city": d.get("city", ""), "state": d.get("state", "")})
    out.sort(key=lambda r: (order.get(r["action"], 9), r["name"]))

    path = "audit_results/triage.csv"
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["action", "reason", "name", "verdict",
                                           "address", "city", "state", "website"])
        w.writeheader(); w.writerows(out)

    from collections import Counter
    counts = Counter(r["action"] for r in out)
    print(f"triage.csv written: {len(out)} flagged listings\n")
    for action in sorted(counts, key=lambda a: order.get(a, 9)):
        print(f"  {action:18} {counts[action]}")
    print(f"\nReport: {os.path.basename(f)}")


if __name__ == "__main__":
    main()
